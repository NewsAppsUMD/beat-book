"""Tests for anchor evidence (names, figures and dates) and for sorting
unsourced claims into facts, analysis and suggestions."""

import citation_matcher as cm
import claim_evidence as ce
from chat_provider import ChatResponse

STORIES = [
    {"article_id": "story-0", "title": "Fire breaks ground",
     "content": ("The Chicago Fire broke ground Tuesday on a $750 million soccer stadium in the South "
                 "Loop at The 78. The stadium is bankrolled by the team's billionaire owner, Joe "
                 "Mansueto, without public money. ") * 2},
    {"article_id": "story-1", "title": "Bears weigh Hammond",
     "content": ("The Bears are weighing a stadium in Hammond. Chicago lawmakers and the Bears "
                 "met again. Fans in Chicago want the Bears to stay. ") * 3},
    {"article_id": "story-2", "title": "Bears fans react",
     "content": ("Chicago Bears fans reacted to the Hammond news on social media. The Bears said "
                 "Chicago remains possible. ") * 3},
    {"article_id": "story-3", "title": "Committee update",
     "content": ("Vikings owner Mark Wilf chairs the NFL stadium committee, and Bears chairman "
                 "George McCaskey is a member of it. The committee meets in May. ") * 2},
]


def _index():
    return {"articles": [{**s, "passages": cm._passage_windows(s["content"])} for s in STORIES]}


def _entries(*claims):
    return {"entries": [{"content": c, "passthrough": False, "kind": "sentence",
                         "provenance": "unsupported", "supports": []} for c in claims], "stats": {}}


def test_names_skip_ordinary_sentence_starters_but_keep_possessives():
    assert ce.names_in("Several fault lines drove the collapse in Springfield.") == ["springfield"]
    assert ce.names_in("Buckner's bill cleared the House committee.") == ["buckner"]
    assert "mansueto" in ce.names_in("Joe Mansueto owns the Chicago Fire.")


def test_names_skip_the_word_that_opens_a_list_items_description():
    assert ce.names_in("- **Mayor Jim Tinaglia** — Skeptical of the state's late proposal.") == ["jim", "tinaglia"]
    assert ce.names_in("- **Late spring** — Self-imposed Bears deadline.") == ["bears"]
    assert ce.names_in("Indiana state Sens. Ryan Mishler and Rick Niemeyer") == ["ryan", "mishler", "rick", "niemeyer"]


def test_fact_with_distinctive_anchors_is_cited_with_highlights():
    e = _entries("Joe Mansueto is bankrolling the Fire's $750 million stadium at The 78 without public money.")
    assert ce.add_anchor_evidence(e, _index()) == 1
    claim = e["entries"][0]
    sup = claim["supports"][0]
    assert claim["provenance"] == "corpus" and sup["match_type"] == "anchors"
    assert sup["article_id"] == "story-0"
    marked = {STORIES[0]["content"][h["char_offset"]:h["char_offset"] + h["char_length"]].lower()
              for h in sup["highlights"]}
    assert {"mansueto", "750"} <= marked


def test_common_names_alone_never_count():
    # "Bears", "Chicago" and "Hammond" are in most stories: not distinctive.
    e = _entries("The Bears' board voted in Chicago to advance the Hammond plans.")
    assert ce.add_anchor_evidence(e, _index()) == 0
    assert e["entries"][0]["provenance"] == "unsupported"


def test_anchors_without_shared_context_do_not_count():
    # Mansueto and 750 appear together, but the claim is about something else.
    e = _entries("Joe Mansueto sold 750 season tickets for a hockey franchise in Denver last winter.")
    assert ce.add_anchor_evidence(e, _index()) == 0


def test_every_anchor_must_be_present():
    e = _entries("Vikings owner Mark Wilf chairs the NFL stadium committee, which met 14 times.")
    assert ce.add_anchor_evidence(e, _index()) == 0          # "14" is in no story
    e = _entries("Vikings owner Mark Wilf chairs the NFL stadium committee; McCaskey sits on it.")
    assert ce.add_anchor_evidence(e, _index()) == 1


class LabelProvider:
    label_model = "fake-label"

    def __init__(self, kinds=None, fail=False):
        self.kinds, self.fail, self.calls = kinds or {}, fail, []

    def create(self, **kw):
        self.calls.append(kw)
        if self.fail:
            raise RuntimeError("model down")
        n = kw["messages"][0]["content"].count("\n") + 1
        labels = [{"id": i, "kind": self.kinds.get(i, "fact")} for i in range(n)]
        return ChatResponse(content=[{"type": "tool_use", "id": "t", "name": "label_claims",
                                      "input": {"labels": labels}}], stop_reason="tool_use")


def test_analysis_and_suggestions_are_relabeled_facts_stay_flagged():
    e = _entries(
        "The Chicago Bears stadium saga is one of the most consequential disputes in a generation.",
        "Illinois's spring legislative session ended May 31, 2026, without a deal.",
        "Get Stacy Davis Gates on the record about the PILOT bill.",
    )
    info = ce.classify_claims(e, LabelProvider({0: "analysis", 2: "suggestion"}))
    assert [x["provenance"] for x in e["entries"]] == ["analysis", "unsupported", "guidance"]
    assert info["analysis"] == 1 and info["fact"] == 1 and info["suggestion"] == 1
    ce.recount(e)
    assert e["stats"]["analysis"] == 1 and e["stats"]["unsupported"] == 1 and e["stats"]["guidance"] == 1


def test_only_unsupported_claims_are_sent_and_failures_keep_them_flagged():
    e = _entries("A claim with no source that names Pritzker.")
    e["entries"].append({"content": "A cited claim.", "passthrough": False, "provenance": "corpus", "supports": [{}]})
    p = LabelProvider(fail=True)
    info = ce.classify_claims(e, p)
    assert "cited claim" not in p.calls[0]["messages"][0]["content"]
    assert e["entries"][0]["provenance"] == "unsupported" and info["errors"]


def test_unsourced_reasons_separate_outside_details_from_scattered_ones():
    e = _entries(
        "The Bears have played in Chicago since 1920 and moved in the Giants-Jets fashion.",
        "Joe Mansueto and the Bears met in Hammond about the stadium.",
        "Dozens of states have such laws.",
    )
    counts = ce.explain_unsourced(e, _index())
    reasons = [x["unsourced_reason"] for x in e["entries"]]
    assert reasons == ["outside_stories", "in_stories", "no_details"]
    assert e["entries"][0]["details_not_in_stories"] == ["Giants-Jets", "1920"]
    assert counts == {"outside_stories": 1, "in_stories": 1, "outcome_not_stated": 0, "no_details": 1}


def test_word_export_explains_unsourced_facts():
    import io
    from docx import Document
    from app import _markdown_to_docx
    entries = [
        {"content": "## Beat Overview", "passthrough": True, "kind": "other", "supports": []},
        {"content": "The Bears have played in Chicago since 1920.", "passthrough": False, "kind": "sentence",
         "provenance": "unsupported", "supports": [], "unsourced_reason": "outside_stories",
         "details_not_in_stories": ["1920"]},
        {"content": "It is one of the most consequential disputes in a generation.", "passthrough": False,
         "kind": "sentence", "provenance": "analysis", "supports": []},
    ]
    stats = {"unsourced_reasons": {"outside_stories": 1, "in_stories": 0, "no_details": 0}}
    doc = Document(io.BytesIO(_markdown_to_docx("", entries, stats)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "About the sourcing" in text and "Claims to check" in text
    assert "1 are analysis or interpretation" in text
    assert "Not in any story: 1920." in text
    assert "writing model's own general knowledge" in text
    # Without stats (older books), no sourcing section is added.
    doc = Document(io.BytesIO(_markdown_to_docx("", entries)))
    assert "About the sourcing" not in "\n".join(p.text for p in doc.paragraphs)


def test_labels_are_read_from_the_shapes_models_return():
    tool = lambda data: ChatResponse(content=[{"type": "tool_use", "id": "t", "name": "label_claims",
                                               "input": data}], stop_reason="tool_use")
    want = {0: "fact", 1: "analysis"}
    assert ce._parse_labels(tool({"labels": [{"id": 0, "kind": "fact"}, {"id": 1, "kind": "analysis"}]})) == want
    assert ce._parse_labels(tool({"labels": [{"id": "0", "label": "Fact"}, {"id": "1", "label": "analysis"}]})) == want
    assert ce._parse_labels(tool({"labels": ["fact", "analysis"]})) == want
    assert ce._parse_labels(tool({"0": "fact", "1": "analysis"})) == want
    assert ce._parse_labels(tool({"labels": {"0": "fact", "1": "analysis"}})) == want
    text = ChatResponse(content=[{"type": "text", "text": 'Here: [{"id": 0, "type": "fact"}, {"id": 1, "type": "analysis"}]'}],
                        stop_reason="end_turn")
    assert ce._parse_labels(text) == want


def test_unreadable_reply_is_recorded():
    class Garbled:
        label_model = "m"

        def create(self, **kw):
            return ChatResponse(content=[{"type": "tool_use", "id": "t", "name": "label_claims",
                                          "input": {"labels": [{"sentence": 0, "verdict": "?"}]}}],
                                stop_reason="tool_use")
    e = _entries("Illinois's spring legislative session ended May 31, 2026, without a deal.")
    info = ce.classify_claims(e, Garbled())
    assert e["entries"][0]["provenance"] == "unsupported"
    assert "verdict" in info["errors"][0]


def test_date_pinned_down_by_a_weekday_is_not_called_outside_the_stories():
    idx = {"articles": [
        {"article_id": "a", "date": "2026-06-05",
         "content": "The Bears' board of directors voted Thursday to advance the Hammond plan, "
                    "then announced it Friday."},
    ]}
    e = _entries("The Bears' board of directors voted June 4 to advance Hammond.",
                 "The Bears' board of directors voted June 2 to advance Hammond.")
    ce.explain_unsourced(e, idx)
    assert e["entries"][0]["unsourced_reason"] == "in_stories"
    assert e["entries"][1]["unsourced_reason"] == "outside_stories"
    assert e["entries"][1]["details_not_in_stories"] == ["June 2"]
