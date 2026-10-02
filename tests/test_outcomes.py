"""Tests for the outcome check: a claim that someone won, lost, was
acquitted or was denied needs a citation that reports it happened."""

import claim_evidence as ce

PREVIEW = ("Longtime Cook County Board President Toni Preckwinkle faces off Tuesday in the Democratic "
           "primary against veteran Chicago Ald. Brendan Reilly. She's running for a fifth, four-year "
           "term. Whoever wins the Democratic primary would likely win in November. Preckwinkle last "
           "won with about 69% of the vote in 2022 against Bob Fioretti. The county employs 20,000 "
           "people and has a budget approved by a board of separately elected commissioners.")
VERDICT = ("A Cook County tax official was acquitted of drunken driving Tuesday after a two-day trial. "
           "A day after her acquittal on a DUI charge, Commissioner Samantha Steele blasted the "
           "prosecutors. Steele lost reelection in March to Liz Nicholson.")
SETTLEMENT = "Court records show the CHA paid Cooper $525,000 last year to settle a 2023 lawsuit."

ARTICLES = [
    {"article_id": "story-0", "title": "Preckwinkle faces Reilly", "content": PREVIEW},
    {"article_id": "story-1", "title": "Steele acquitted", "content": VERDICT},
    {"article_id": "story-2", "title": "Cooper settlement", "content": SETTLEMENT},
]


def _support(aid, start=0, length=None):
    content = next(a["content"] for a in ARTICLES if a["article_id"] == aid)
    length = len(content) if length is None else length
    return {"article_id": aid, "passage_text": content[start:start + length],
            "passage_offset": start, "passage_length": length, "similarity": 0.8}


def _cited(claim, *supports):
    return {"entries": [{"content": claim, "passthrough": False, "kind": "sentence",
                         "provenance": "corpus", "supports": list(supports)}], "stats": {}}


def test_outcomes_found_in_claims_skip_quotes_and_hypotheticals():
    assert ce.outcomes_in("Preckwinkle turned back Reilly.") == [("election", "turned back")]
    assert ce.outcomes_in("Steele lost, then was acquitted in May.") == [
        ("election", "lost"), ("acquittal", "acquitted")]
    assert ce.outcomes_in("If Preckwinkle wins, it will be her last term.") == []
    assert ce.outcomes_in('Brewer said the mayor "fired" nobody.') == []
    assert ce.outcomes_in("The beat covers separately elected offices.") == []
    # Conceding a point isn't conceding a race.
    assert ce.outcomes_in("Buckner conceded the point, calling it a first step.") == []
    assert ce.outcomes_in("Reilly conceded to Preckwinkle.") == [("election", "conceded")]


def test_preview_story_does_not_support_a_result():
    # The preview mentions a past win (2022), a hypothetical ("Whoever
    # wins") and "elected" as an adjective. None reports this year's result.
    e = _cited("Toni Preckwinkle turned back Ald. Brendan Reilly in her bid for a fifth term.",
               _support("story-0"))
    assert ce.check_outcomes(e, {"articles": ARTICLES}) == 1
    claim = e["entries"][0]
    assert claim["provenance"] == "unsupported" and claim["supports"] == []
    assert claim["outcome_not_stated"] == ["turned back"]
    ce.explain_unsourced(e, {"articles": ARTICLES})
    assert claim["unsourced_reason"] == "outcome_not_stated"


def test_citation_that_reports_the_outcome_is_kept():
    e = _cited("Steele was acquitted of drunken driving.", _support("story-1"))
    assert ce.check_outcomes(e, {"articles": ARTICLES}) == 0
    assert e["entries"][0]["provenance"] == "corpus"


def test_year_describing_a_thing_does_not_disqualify_the_sentence():
    # "a 2023 lawsuit" isn't the settlement's date.
    e = _cited("The CHA settled Cooper's case for $525,000 in 2025.", _support("story-2"))
    assert ce.check_outcomes(e, {"articles": ARTICLES}) == 0


def test_missing_outcome_is_found_elsewhere_in_the_stories():
    # Matched only to the preview, but another story reports both outcomes
    # in sentences naming Steele (and Nicholson).
    e = _cited("Samantha Steele lost to Liz Nicholson, then was acquitted.", _support("story-0"))
    assert ce.check_outcomes(e, {"articles": ARTICLES}) == 0
    claim = e["entries"][0]
    assert claim["provenance"] == "corpus"
    found = [s for s in claim["supports"] if s["match_type"] == "outcome"]
    assert {s["article_id"] for s in found} == {"story-1"}
    marked = [VERDICT[h["char_offset"]:h["char_offset"] + h["char_length"]]
              for s in found for h in s["highlights"]]
    assert any("acquittal" in m for m in marked) and any("lost reelection" in m for m in marked)


def test_anchor_evidence_requires_the_outcome():
    import citation_matcher as cm
    # Filler stories, so Fioretti is in few enough of them to be distinctive.
    filler = [{"article_id": f"filler-{i}", "content": "The county board met and adjourned. " * 5}
              for i in range(4)]
    idx = {"articles": [{**a, "passages": cm._passage_windows(a["content"])} for a in ARTICLES + filler]}

    def entries(claim):
        return {"entries": [{"content": claim, "passthrough": False, "kind": "sentence",
                             "provenance": "unsupported", "supports": []}], "stats": {}}
    # Same anchors (Preckwinkle, Fioretti, 69): cited without an outcome,
    # refused with one the story never reports.
    plain = entries("Preckwinkle took about 69% of the vote against Bob Fioretti.")
    assert ce.add_anchor_evidence(plain, idx) == 1
    outcome = entries("Preckwinkle was acquitted after taking about 69% of the vote against Bob Fioretti.")
    assert ce.add_anchor_evidence(outcome, idx) == 0


def test_word_export_names_the_outcome():
    import io
    from docx import Document
    from app import _markdown_to_docx
    entries = [{"content": "Preckwinkle turned back Reilly.", "passthrough": False, "kind": "sentence",
                "provenance": "unsupported", "supports": [], "unsourced_reason": "outcome_not_stated",
                "outcome_not_stated": ["turned back"]}]
    stats = {"unsourced_reasons": {"outside_stories": 0, "in_stories": 0, "outcome_not_stated": 1,
                                   "no_details": 0}}
    text = "\n".join(p.text for p in Document(io.BytesIO(_markdown_to_docx("", entries, stats))).paragraphs)
    assert "1 state an outcome that no matching passage reports" in text
    assert "Outcome it states: “turned back”." in text


def test_claim_naming_no_one_can_use_a_cited_story():
    # A sub-bullet under "**Debra Parker**" says "she". The cited story
    # reports the outcome in a sentence sharing the claim's key words.
    story = {"article_id": "p", "title": "Parker drops suit",
             "content": ("The commissioner has served since 2018. " + "The board met again this week. " * 30
                         + "Longtime CHA Board Commissioner Debra Parker dropped her lawsuit Friday against "
                         "the housing authority over her voucher.")}
    e = _cited("She dropped her lawsuit over the voucher on March 27.",
               {"article_id": "p", "passage_text": story["content"][:40], "passage_offset": 0,
                "passage_length": 40, "similarity": 0.8})
    assert ce.check_outcomes(e, {"articles": [story]}) == 0
    assert e["entries"][0]["supports"][0]["match_type"] == "outcome"
    # Without shared key words, a nameless claim isn't rescued.
    e = _cited("She dropped the case.", {"article_id": "p", "passage_text": story["content"][:40],
                                   "passage_offset": 0, "passage_length": 40, "similarity": 0.8})
    assert ce.check_outcomes(e, {"articles": [story]}) == 1


def _story(text, aid="s"):
    return {"articles": [{"article_id": aid, "title": "Primary results", "date": "2026-03-18", "content": text}]}


def _claim(text, provenance="corpus"):
    return {"entries": [{"content": text, "passthrough": False, "kind": "sentence", "provenance": provenance,
                         "supports": [{"article_id": "s", "passage_text": "x"}]}], "stats": {}}


def test_a_reversed_outcome_is_contradicted_with_the_story_sentence():
    e = _claim("Despite these controversies, Steele won the primary in May 2026 after an acquittal.")
    idx = _story("Turnout was low. Steele lost reelection in March to longtime Democratic operative Liz Nicholson.")
    assert ce.check_contradictions(e, idx) == 1
    claim = e["entries"][0]
    assert claim["provenance"] == "unsupported" and claim["supports"] == []
    assert claim["contradicted_by"]["sentence"].startswith("Steele lost reelection")
    ce.explain_unsourced(e, idx)
    assert claim["unsourced_reason"] == "contradicted"


def test_object_of_a_defeat_and_agreement_are_read():
    assert ce.check_contradictions(_claim("Steele won the District 2 primary."),
                                   _story("Liz Nicholson defeated Steele in the District 2 primary.")) == 1
    # A story that agrees anywhere keeps the claim.
    assert ce.check_contradictions(_claim("Nicholson won the District 2 primary."),
                                   _story("Liz Nicholson defeated Steele in the District 2 primary.")) == 0


def test_other_years_and_hypotheticals_do_not_contradict():
    # Her first win, years earlier, isn't about this primary.
    assert ce.check_contradictions(_claim("Steele won her first race in 2020."),
                                   _story("Steele lost reelection in 2026 to Liz Nicholson.")) == 0
    assert ce.check_contradictions(_claim("Steele won the primary."),
                                   _story("If Steele loses, Nicholson takes the seat.")) == 0


def test_election_evidence_must_name_the_claims_person():
    # The cited passage reports someone else's election, and a trial verdict.
    passage = ("She said she was not surprised by the backlash after she unseated an incumbent. "
               "Her attorney said that justice prevailed. Larry Rogers Jr., won his sixth term in 2024.")
    story = {"article_id": "p", "title": "Steele", "content": passage}
    e = _cited("Samantha Steele won the primary.",
               {"article_id": "p", "passage_text": passage, "passage_offset": 0, "passage_length": len(passage),
                "similarity": 0.8})
    assert ce.check_outcomes(e, {"articles": [story]}) == 1
    # A title's period doesn't split the sentence that names the person.
    e = _cited("Larry Rogers Jr. won his sixth term in 2024.",
               {"article_id": "p", "passage_text": passage, "passage_offset": 0, "passage_length": len(passage),
                "similarity": 0.8})
    assert ce.check_outcomes(e, {"articles": [story]}) == 0


def test_word_export_quotes_the_contradicting_sentence():
    import io
    from docx import Document
    from app import _markdown_to_docx
    entries = [{"content": "Steele won the primary.", "passthrough": False, "kind": "sentence",
                "provenance": "unsupported", "supports": [], "unsourced_reason": "contradicted",
                "contradicted_by": {"article_title": "Primary results", "article_date": "2026-03-18",
                                    "sentence": "Steele lost reelection to Liz Nicholson."}}]
    stats = {"unsourced_reasons": {"contradicted": 1}}
    text = "\n".join(p.text for p in Document(io.BytesIO(_markdown_to_docx("", entries, stats))).paragraphs)
    assert "1 is contradicted by a story" in text
    assert "Primary results, 2026-03-18: “Steele lost reelection to Liz Nicholson.”" in text
